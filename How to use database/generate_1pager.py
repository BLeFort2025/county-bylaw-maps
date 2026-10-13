"""
generate_1pager.py — Generates a polished, professional 1-page plain-language overview
of the OFA Municipal Bylaw Database in Word (.docx) format.
"""
import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

HERE = os.path.dirname(os.path.abspath(__file__))

TARGET_FILES = [
    os.path.join(HERE, "OFA Bylaw Database - 1-Page Plain Language Overview.docx"),
    os.path.join(HERE, "Municipal Bylaw Database - Plain Language Overview.docx")
]

def build_document():
    doc = Document()

    # ── Page setup: Margins tuned for exact 1-page / 2-page print readiness ──
    for section in doc.sections:
        section.top_margin = Cm(1.2)
        section.bottom_margin = Cm(1.2)
        section.left_margin = Cm(1.5)
        section.right_margin = Cm(1.5)

    # ── Base Font ──
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Calibri'
    font.size = Pt(9.5)
    font.color.rgb = RGBColor(0x22, 0x22, 0x22)

    DARK_GREEN = RGBColor(0x1B, 0x5E, 0x20)
    CHARCOAL = RGBColor(0x33, 0x33, 0x33)

    def set_cell_background(cell, fill_hex):
        tcPr = cell._tc.get_or_add_tcPr()
        tcPr.append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>'))

    def set_cell_margins(cell, top=80, bottom=80, left=120, right=120):
        tcPr = cell._tc.get_or_add_tcPr()
        tcMar = OxmlElement('w:tcMar')
        for m, val in [('w:top', top), ('w:bottom', bottom), ('w:left', left), ('w:right', right)]:
            node = OxmlElement(m)
            node.set(qn('w:w'), str(val))
            node.set(qn('w:type'), 'dxa')
            tcMar.append(node)
        tcPr.append(tcMar)

    # ── Header ──
    p_head = doc.add_paragraph()
    p_head.paragraph_format.space_before = Pt(0)
    p_head.paragraph_format.space_after = Pt(1)
    run_org = p_head.add_run("ONTARIO FEDERATION OF AGRICULTURE  |  MUNICIPAL ADVOCACY & POLICY")
    run_org.font.size = Pt(8.5)
    run_org.font.bold = True
    run_org.font.color.rgb = DARK_GREEN

    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(2)
    run_title = p_title.add_run("Municipal Bylaw Database & Intelligence System")
    run_title.font.size = Pt(17)
    run_title.font.bold = True
    run_title.font.color.rgb = DARK_GREEN

    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_before = Pt(0)
    p_sub.paragraph_format.space_after = Pt(4)
    run_sub = p_sub.add_run("A Plain-Language Overview for OFA Leaders, County Federations & Farmers")
    run_sub.font.size = Pt(10.5)
    run_sub.font.bold = True
    run_sub.font.color.rgb = CHARCOAL

    # XML Divider Line under title
    p_sub_pr = p_sub._p.get_or_add_pPr()
    pBdr = parse_xml(r'<w:pBdr xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:bottom w:val="single" w:sz="12" w:space="4" w:color="1B5E20"/></w:pBdr>')
    p_sub_pr.append(pBdr)

    # ── 1. What Is It & Why Does It Matter? ──
    p_sec1 = doc.add_paragraph()
    p_sec1.paragraph_format.space_before = Pt(4)
    p_sec1.paragraph_format.space_after = Pt(2)
    r_sec1 = p_sec1.add_run("1. What Is This Tool & Why Does It Matter?")
    r_sec1.font.size = Pt(11)
    r_sec1.font.bold = True
    r_sec1.font.color.rgb = DARK_GREEN

    p_desc1 = doc.add_paragraph()
    p_desc1.paragraph_format.space_before = Pt(0)
    p_desc1.paragraph_format.space_after = Pt(4)
    p_desc1.paragraph_format.line_spacing = 1.15
    r1 = p_desc1.add_run(
        "Ontario has 444 municipal governments (townships, towns, cities, and counties). Every month, local councils enact or update bylaws that directly impact everyday farm operations—from adding tens of thousands of dollars in fees on new barns, to restricting working guardian dogs, regulating fence lines, or imposing runoff charges. When rules change, farmers risk unfair costs unless an explicit agricultural exemption is in place.\n"
        "This tool is OFA's province-wide digital map and early warning system. It tracks exactly which municipalities protect farms, flags upcoming bylaw expiries, and alerts OFA staff and local directors the moment a local council starts debating agricultural bylaws."
    )
    r1.font.size = Pt(9.5)

    # ── 2. The 7 Key Bylaw Categories Tracked ──
    p_sec2 = doc.add_paragraph()
    p_sec2.paragraph_format.space_before = Pt(4)
    p_sec2.paragraph_format.space_after = Pt(3)
    r_sec2 = p_sec2.add_run("2. The 7 Key Bylaw Categories Tracked Across All 444 Municipalities")
    r_sec2.font.size = Pt(11)
    r_sec2.font.bold = True
    r_sec2.font.color.rgb = DARK_GREEN

    table_cats = doc.add_table(rows=8, cols=2)
    table_cats.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_cats.autofit = False

    cat_data = [
        ("Development Charges (DC)", "Building fees levied on new construction. Tracks whether farm buildings (barns, silos, packhouses) and on-farm worker housing are exempt from municipal growth fees."),
        ("Stormwater Fees", "Runoff user fees based on roof/impervious area. Tracks whether agricultural properties and greenhouses are exempt from urban runoff charges."),
        ("Site Alteration & Fill", "Rules on moving dirt, grading land, and accepting clean fill. Ensures normal farming operations and field drainage improvements are protected."),
        ("Livestock Guardian Dogs", "Dog control & kennel rules. Protects guardian and herding dogs as essential working farm animals—exempting them from pet limits, barking bans, and license fees."),
        ("Tree / Forest Conservation", "Tree-cutting bylaws. Tracks whether farmers maintain the right to clear fence lines, remove hazard trees, and manage woodlots for normal farming."),
        ("Backyard Chickens", "Tracks urban hen bylaws, flock limits (usually 3-6 hens), rooster bans, and biosecurity/welfare standards across residential zones."),
        ("Fence Bylaws", "Division fences and line agreements. Tracks whether local rules allow electrified or security fencing for livestock and honor provincial Line Fences Act cost-sharing.")
    ]

    hdr_cells = table_cats.rows[0].cells
    hdr_cells[0].text = "Bylaw Category"
    hdr_cells[1].text = "Why It Matters to Farmers & What We Track"
    for c in hdr_cells:
        set_cell_background(c, "1B5E20")
        set_cell_margins(c, top=70, bottom=70, left=120, right=120)
        for p in c.paragraphs:
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            for run in p.runs:
                run.font.bold = True
                run.font.size = Pt(9)
                run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    for i, (name, desc) in enumerate(cat_data):
        row_cells = table_cats.rows[i+1].cells
        row_cells[0].text = name
        row_cells[1].text = desc
        bg = "F9FBF9" if i % 2 == 0 else "FFFFFF"
        for idx, c in enumerate(row_cells):
            set_cell_background(c, bg)
            set_cell_margins(c, top=50, bottom=50, left=120, right=120)
            for p in c.paragraphs:
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                for run in p.runs:
                    run.font.size = Pt(9)
                    if idx == 0:
                        run.font.bold = True
                        run.font.color.rgb = DARK_GREEN

    for row in table_cats.rows:
        row.cells[0].width = Inches(2.2)
        row.cells[1].width = Inches(5.1)

    # ── 3. How Anyone Can Use It: Traffic Light System ──
    p_sec3 = doc.add_paragraph()
    p_sec3.paragraph_format.space_before = Pt(5)
    p_sec3.paragraph_format.space_after = Pt(2)
    r_sec3 = p_sec3.add_run("3. The Interactive Map: A Simple 'Traffic Light' Visual")
    r_sec3.font.size = Pt(11)
    r_sec3.font.bold = True
    r_sec3.font.color.rgb = DARK_GREEN

    p_desc3 = doc.add_paragraph()
    p_desc3.paragraph_format.space_before = Pt(0)
    p_desc3.paragraph_format.space_after = Pt(3)
    p_desc3.paragraph_format.line_spacing = 1.15
    r3 = p_desc3.add_run(
        "You do not need any technical training to use the system. When you open the website in your browser, you choose an issue (e.g., 'Development Charges') and the entire province of Ontario lights up clearly:\n"
        "• GREEN: Farm Exemption in Place - Local farmers are protected.\n"
        "• RED: No Farm Exemption - Agriculture is charged or restricted (immediate priority for OFA advocacy).\n"
        "• GREY: Not Applicable / No Bylaw - The municipality has not enacted a bylaw for this topic.\n"
        "• BORDER OUTLINES: An Orange Border means a bylaw is expiring within 18 months. A Red-Orange Border means our scanner caught active council discussions in recent meeting minutes."
    )
    r3.font.size = Pt(9.5)

    # ── 4. Early Warning & 1-Click Action ──
    p_sec4 = doc.add_paragraph()
    p_sec4.paragraph_format.space_before = Pt(4)
    p_sec4.paragraph_format.space_after = Pt(2)
    r_sec4 = p_sec4.add_run("4. Automated Early Warning & 1-Click Advocacy Outreach")
    r_sec4.font.size = Pt(11)
    r_sec4.font.bold = True
    r_sec4.font.color.rgb = DARK_GREEN

    p_desc4 = doc.add_paragraph()
    p_desc4.paragraph_format.space_before = Pt(0)
    p_desc4.paragraph_format.space_after = Pt(4)
    p_desc4.paragraph_format.line_spacing = 1.15
    r4 = p_desc4.add_run(
        "1. Weekly Council Scanner: The system automatically reads council meeting agendas and minutes across hundreds of municipal websites every Sunday night. If a town starts considering a bylaw review, it flags the town automatically.\n"
        "2. One-Click Letters to Municipal Clerks: When you click any municipality on the map that lacks an agricultural exemption, the tool can instantly generate a personalized, pre-formatted Microsoft Word advocacy letter. It automatically looks up the local Municipal Clerk's email address from our 7,000+ contact directory, populates OFA's official policy talking points, and even drafts an email in Outlook ready to send."
    )
    r4.font.size = Pt(9.5)

    # ── Bottom Summary Box ──
    table_box = doc.add_table(rows=1, cols=1)
    table_box.alignment = WD_TABLE_ALIGNMENT.CENTER
    box_cell = table_box.rows[0].cells[0]
    set_cell_background(box_cell, "E8F5E9")
    set_cell_margins(box_cell, top=80, bottom=80, left=150, right=150)
    box_p = box_cell.paragraphs[0]
    box_p.paragraph_format.space_before = Pt(0)
    box_p.paragraph_format.space_after = Pt(0)
    r_box = box_p.add_run(
        "Quick Takeaway: This tool turns 444 municipal websites and thousands of pages of confusing legal bylaws into a simple, 1-click map. It ensures OFA directors and staff know exactly where farmers are protected, alerts us before councils vote on bad policy, and equips local leaders to take action immediately."
    )
    r_box.font.bold = True
    r_box.font.size = Pt(9.5)
    r_box.font.color.rgb = DARK_GREEN

    return doc

if __name__ == "__main__":
    doc = build_document()
    for path in TARGET_FILES:
        doc.save(path)
        print(f"Saved: {path}")
